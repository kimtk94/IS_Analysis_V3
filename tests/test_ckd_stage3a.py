#!/usr/bin/env python3
import importlib.util
import unittest
from pathlib import Path

P = Path(__file__).resolve().parents[1] / "scripts" / "run_ckd_stage3a_kidney_evidence.py"
spec = importlib.util.spec_from_file_location("m", P)
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)


class Stage3ATests(unittest.TestCase):
    def test_figshare_file_id(self):
        self.assertEqual(
            m.figshare_file_id("https://figshare.com/ndownloader/files/33957947"),
            33957947,
        )
        self.assertIsNone(m.figshare_file_id("https://example.org/x"))

    def test_validate_gzip(self):
        import gzip, tempfile
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.txt.gz"
            with gzip.open(p, "wb") as fh:
                fh.write(b"hello\n")
            self.assertTrue(m.validate_gzip(p))
            p.write_bytes(b"<html>blocked</html>")
            self.assertFalse(m.validate_gzip(p))

    def test_main_text_candidate_evidence(self):
        self.assertAlmostEqual(
            m.HIROHAMA_MAIN_TEXT["ACP1"]["kidney_egfr_pph4"], 0.999
        )
        self.assertAlmostEqual(
            m.HIROHAMA_MAIN_TEXT["GSTA1"]["kidney_egfr_pph4"], 0.999
        )
        self.assertAlmostEqual(
            m.HIROHAMA_MAIN_TEXT["INHBC"]["kidney_egfr_pph4"], 0.875
        )

    def test_valid_xlsx_guard(self):
        import tempfile, zipfile
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.xlsx"
            p.write_bytes(b"<html>not xlsx</html>")
            self.assertFalse(m.is_valid_xlsx(p))
            with zipfile.ZipFile(p, "w") as z:
                z.writestr("[Content_Types].xml", "<Types/>")
                z.writestr("xl/worksheets/sheet1.xml", "<worksheet/>")
                z.writestr("padding.bin", b"0" * 120000)
            self.assertTrue(m.is_valid_xlsx(p))

    def test_workbook_supplement_heuristic(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "supp.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Supplementary Table 1"
            ws["A1"] = "Supplementary Table 1"
            for i in range(2, 5):
                wb.create_sheet(f"Supplementary Table {i}")
            wb.save(p)
            # Naming/text evidence is enough even when the synthetic file is small.
            self.assertTrue(m.workbook_looks_like_supp_tables(p))

    def test_gene_match_exact_and_multivalue(self):
        genes = {"F12", "GSTA3", "ACP1"}
        self.assertEqual(m.exact_gene_in_cell("F12", genes), "F12")
        self.assertEqual(
            m.exact_gene_in_cell("ABC;GSTA3;XYZ", genes), "GSTA3"
        )
        self.assertIsNone(m.exact_gene_in_cell("F12A1", genes))

    def test_table_number_and_evidence(self):
        self.assertEqual(
            m.sheet_table_number("Supplementary Table 3"), 3
        )
        self.assertEqual(
            m.evidence_type_from_table(3),
            "kidney_cis_pqtl_significant",
        )
        self.assertEqual(
            m.evidence_type_from_table(13),
            "kidney_cis_eqtl_significant_same_study",
        )
        self.assertEqual(
            m.evidence_type_from_table(16),
            "kidney_qtl_egfr_colocalization",
        )

    def test_stage2_classification(self):
        self.assertEqual(
            m.classify_stage2({
                "comparison_status": "ABF_and_SuSiE_shared_signal"
            }),
            "core_shared_signal",
        )
        self.assertEqual(
            m.classify_stage2({
                "comparison_status": "SuSiE_nonconverged_or_failed"
            }),
            "susie_unresolved",
        )
        self.assertEqual(
            m.classify_stage2({
                "comparison_status": "ABF_not_confirmed_by_SuSiE",
                "max_PP.H4": "0.77",
            }),
            "abf_supported_susie_weaker",
        )
        self.assertEqual(
            m.classify_stage2({
                "comparison_status": "ABF_not_confirmed_by_SuSiE",
                "max_PP.H4": "3e-14",
            }),
            "abf_susie_discordant",
        )


if __name__ == "__main__":
    unittest.main()
