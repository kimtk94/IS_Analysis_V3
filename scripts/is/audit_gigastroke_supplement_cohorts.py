#!/usr/bin/env python3
"""Public supplementary-table source audit; do not infer sample overlap."""
import hashlib,json,re,sys,urllib.parse,urllib.request,zipfile
from html.parser import HTMLParser
from io import BytesIO
from openpyxl import load_workbook
ARTICLE="https://pmc.ncbi.nlm.nih.gov/articles/PMC9524349/"
class Anchors(HTMLParser):
 def __init__(self):super().__init__();self.items=[];self.href=None;self.label=[]
 def handle_starttag(self,tag,attrs):
  if tag=="a":self.href=dict(attrs).get("href");self.label=[]
 def handle_data(self,data):
  if self.href is not None:self.label.append(data)
 def handle_endtag(self,tag):
  if tag=="a" and self.href is not None:self.items.append(("".join(self.label).strip(),self.href));self.href=None
def get(url):
 req=urllib.request.Request(url,headers={"User-Agent":"IS-Research-Audit/1.0 (public supplemental retrieval; research)"})
 with urllib.request.urlopen(req,timeout=50) as r:
  content=r.read(12_000_001);final=r.geturl()
 if len(content)>12_000_000:raise ValueError("SOURCE_TOO_LARGE")
 if urllib.parse.urlsplit(final).hostname not in ("pmc.ncbi.nlm.nih.gov","www.ncbi.nlm.nih.gov","static-content.springer.com","media.springernature.com","static-content.springer.com"):
  raise ValueError("UNTRUSTED_DOWNLOAD_HOST "+str(final))
 return content,final
def main():
 page,_=get(ARTICLE);parser=Anchors();parser.feed(page.decode("utf-8","replace"))
 refs=[(label,urllib.parse.urljoin(ARTICLE,href)) for label,href in parser.items if ("supplementary tables" in label.lower() and ("xlsx" in href.lower() or "supp" in href.lower() or "bin/" in href.lower()))]
 print("SUPP_CANDIDATE_LINKS="+json.dumps(refs[:12]),flush=True)
 if not refs:raise RuntimeError("NO_SUPPLEMENTAL_TABLE_LINK")
 fallbacks=[refs[-1][1],"https://static-content.springer.com/esm/art%3A10.1038%2Fs41586-022-05165-3/MediaObjects/41586_2022_5165_MOESM4_ESM.xlsx"]
 attempts=[]
 data=final=None
 for url in fallbacks:
  try:
   chunk,resolved=get(url)
   if zipfile.is_zipfile(BytesIO(chunk)):
    data,final=chunk,resolved
    break
   attempts.append({"url":url,"resolved":resolved,"length":len(chunk),"prefix":repr(chunk[:100])})
  except Exception as e: attempts.append({"url":url,"error":repr(e)})
 if data is None:
  print("SUPP_DOWNLOAD_ATTEMPTS="+json.dumps(attempts),flush=True)
  raise RuntimeError("NO_VALID_OFFICIAL_XLSX")
 book=load_workbook(BytesIO(data),read_only=True,data_only=True)
 names=book.sheetnames
 print("SUPP_XLSX_PROVENANCE="+json.dumps({"url":final,"sha256":hashlib.sha256(data).hexdigest(),"size":len(data),"sheet_count":len(names),"sheets":names[:60]}),flush=True)
 for sn in names:
  if not re.search(r"(^|\\D)(1|2)(\\D|$)",sn) and not re.search("table\\s*[12]\\b",sn,re.I):continue
  ws=book[sn];values=list(ws.iter_rows(max_row=400,max_col=45,values_only=True))
  hits=[]
  current_ancestry=""
  for idx,row in enumerate(values,1):
   if idx>4 and row[0] is not None and str(row[0]).strip() in ("EUR","EAS","AFR","SAS","HIS","AMR"):
    current_ancestry=str(row[0]).strip()
   study=str(row[1]).strip() if len(row)>1 and row[1] is not None else ""
   if idx<=4 or not study:continue
   if current_ancestry=="EAS" or any(term in study.casefold() for term in ("bbj","biobank japan","japan","tohoku")):
    hits.append({"row":idx,"ancestry":current_ancestry,"study":study,"cells":[str(x)[:200] if x is not None else "" for x in row[:23]]})
  print("SUPP_TABLE_EXTRACT="+json.dumps({"sheet":sn,"dimensions":[ws.max_row,ws.max_column],"header_rows":[[str(x)[:180] if x is not None else "" for x in row[:18]] for row in values[:8]],"eas_and_japan_rows":hits[:100]},ensure_ascii=False),flush=True)
 print("SUPP_EVIDENCE_AUDIT=COMPLETE_DESCRIPTIVE_ONLY",flush=True)
if __name__=="__main__":
 try:main()
 except Exception as exc:print("SUPP_AUDIT_ERROR="+repr(exc),file=sys.stderr,flush=True);sys.exit(1)
