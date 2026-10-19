#!/usr/bin/env python3
"""CI-only public data source verification; no server, private data or causal claims."""
import datetime as dt,hashlib,json,os,pathlib,urllib.request
from io import BytesIO
from openpyxl import load_workbook
from discover_public_bbj_is_gwas import source as discover_bbj
URL="https://static-content.springer.com/esm/art%3A10.1038%2Fs41586-022-05165-3/MediaObjects/41586_2022_5165_MOESM4_ESM.xlsx"
PIN="7c688f864c73ec2e9624bde9f53ac48a9880a6a08e59f5e1266690e243618ba4"
def analyze(data):
 digest=hashlib.sha256(data).hexdigest()
 if digest!=PIN:raise ValueError("SUPPLEMENT_SHA256_DRIFT")
 w=load_workbook(BytesIO(data),read_only=True,data_only=True)
 sheet=w["Suppl.Table1"]
 selected=[]
 for rowno,row in enumerate(sheet.iter_rows(values_only=True),1):
  if str(row[1]).strip().upper().startswith("BIOBANK_JAPAN"):
   selected.append(dict(excel_row=rowno,study=str(row[1]),AIS_cases=row[10],controls=row[16-1]))
 if len(selected)!=1:raise ValueError("BBJ_ROW_COUNT_DRIFT")
 row=selected[0]
 if int(row["AIS_cases"])!=17651:raise ValueError("BBJ_AIS_DRIFT")
 return {"source_url":URL,"sha256":digest,"bbj_eas_discovery":row}
def main():
 output=pathlib.Path("ci_public_gwas_evidence.json")
 state={"schema":"is-ci-public-source-gate-v1","generated_at":dt.datetime.now(dt.timezone.utc).isoformat(),
        "source_commit":os.getenv("GITHUB_SHA",""),"mode":"CI_ONLY_NO_SSH",
        "independent_replication":"NOT_ESTABLISHED","live_writes":False}
 try:
  request=urllib.request.Request(URL,headers={"User-Agent":"IS-Public-Research-Harness/1.0"})
  with urllib.request.urlopen(request,timeout=40) as f:raw=f.read(10_000_001)
  if len(raw)>10_000_000:raise ValueError("SUPPLEMENT_TOO_LARGE")
  state["gigastroke"]=analyze(raw)
  state["gigastroke_status"]="PASS"
 except (OSError,TimeoutError) as exc:
  state["gigastroke_status"]="SOURCE_UNAVAILABLE"
  state["gigastroke_error"]=type(exc).__name__
 except Exception as exc:
  state["gigastroke_status"]="EVIDENCE_REVIEW_REQUIRED"
  state["gigastroke_error"]=str(exc)[:250]
 try:
  matches=discover_bbj()
  state["bbj_2021"]={"status":"DISCOVERED" if matches else "NO_MATCH_IN_INDEX","matches":matches,
                      "file_downloaded":False,"cohort_independent":False}
 except Exception as exc:
  state["bbj_2021"]={"status":"SOURCE_UNAVAILABLE","error":type(exc).__name__}
 state["status"]="PASS" if state["gigastroke_status"]=="PASS" and state["bbj_2021"]["status"]=="DISCOVERED" else "REVIEW_OR_SOURCE_PENDING"
 output.write_text(json.dumps(state,indent=2,ensure_ascii=False)+"\n")
 print("CI_PUBLIC_SOURCE_STATUS="+state["status"],flush=True)
 print("CI_PUBLIC_GIGASTROKE="+state["gigastroke_status"],flush=True)
 print("CI_PUBLIC_BBJ="+state["bbj_2021"]["status"],flush=True)
 if state["gigastroke_status"]=="EVIDENCE_REVIEW_REQUIRED":raise SystemExit(1)
if __name__=="__main__":main()
