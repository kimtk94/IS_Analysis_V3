#!/usr/bin/env python3
"""P5j: official IOVS 2026 TM 48h signed DEG replication crosswalk.

Inputs: official unmodified Europe PMC supplementary ZIP, verified P5b and P5d
aggregate-only GitHub Actions artifacts, P5g-derived shortlisted genes.
No donor profiles; do not equate absence of a published DEG with novel biology.
"""
import argparse
import csv
import gzip
import io
import json
import math
import statistics
import urllib.request
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

SOURCE="https://www.ebi.ac.uk/europepmc/webservices/rest/PMC12882097/supplementaryFiles"
OFFICIAL_NAME="iovs-67-2-13_s003.xlsx"
TM_ALL="SuppTable4_TM2Days_All_DEGs"
TM_HIGH="SuppTable3_TM2Days_HighConfDEGs"

def norm_gene(v):
    return str(v or "").strip().split(".")[0].upper()

def finite(v):
    try:
        f=float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None

def download_zip(path):
    request=urllib.request.Request(SOURCE,headers={"User-Agent":"ocular-p5j-replication-audit/1.0"})
    with urllib.request.urlopen(request, timeout=120) as resp, open(path,"wb") as out:
        limit=70*1024*1024
        transferred=0
        while True:
            b=resp.read(1024*1024)
            if not b:break
            transferred+=len(b)
            if transferred>limit:raise ValueError("SUPPLEMENT_ZIP_GT_70MB")
            out.write(b)
    return transferred

def parse_official(path):
    with zipfile.ZipFile(path) as archive:
        matches=[name for name in archive.namelist()
                 if name.rsplit("/",1)[-1]==OFFICIAL_NAME]
        if len(matches)!=1:raise ValueError("OFFICIAL_WORKBOOK_NOT_FOUND")
        content=archive.read(matches[0])
    wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    if not {TM_ALL,TM_HIGH}<=set(wb.sheetnames):
        raise ValueError("TM_SHEETS_MISSING "+str(wb.sheetnames))
    result={}
    for name in [TM_ALL,TM_HIGH]:
        sheet=wb[name]
        it=sheet.iter_rows(values_only=True)
        header=[str(x).strip() if x is not None else "" for x in next(it)]
        need={"ID","symbol","padj"}
        if name==TM_ALL:need|={"log2FoldChange"}
        else:need|={"log2FoldChange_TM_2Days","padj_TM_2Days"}
        if name==TM_HIGH:
            need.remove("padj")
        if not need<=set(header):
            raise ValueError(name+" MISSING_COLUMNS "+str(need-set(header)))
        col={v:i for i,v in enumerate(header)}
        records={}
        for row in it:
            if not row or not row[col["ID"]]:continue
            gid=norm_gene(row[col["ID"]])
            lfc=finite(row[col["log2FoldChange" if name==TM_ALL else "log2FoldChange_TM_2Days"]])
            fdr=finite(row[col["padj" if name==TM_ALL else "padj_TM_2Days"]])
            if not gid.startswith("ENSG") or lfc is None:
                raise ValueError(name+" INVALID_GENE_ID_OR_LFC")
            if fdr is not None and not 0<=fdr<=1:
                raise ValueError(name+" INVALID_ADJUSTED_P")
            if gid in records:raise ValueError(name+" DUPLICATE_GENE_ID "+gid)
            records[gid]={"gene_id":gid,"symbol":str(row[col["symbol"]] or "").strip(),
                          "log2fc":lfc,"fdr":fdr}
        result[name]=records
    wb.close()
    if not 13000<=len(result[TM_ALL])<=14000:
        raise ValueError("UNEXPECTED_OFFICIAL_TM_ALL_ROWS")
    if not 800<=len(result[TM_HIGH])<=900:
        raise ValueError("UNEXPECTED_OFFICIAL_TM_HIGH_ROWS")
    if not set(result[TM_HIGH])<=set(result[TM_ALL]):
        raise ValueError("HIGH_CONF_NOT_SUBSET_OF_ALL")
    return result

def load_primary(path):
    with open(path,"rb") as raw, gzip.GzipFile(fileobj=raw) as gz:
        it=csv.DictReader(io.TextIOWrapper(gz,encoding="utf8"),delimiter="\t")
        if not {"gene_id","gene_symbol","logFC","FDR","PValue"}<=set(it.fieldnames or []):
            raise ValueError("PRIMARY_FIELDS_MISSING")
        out={}
        for r in it:
            gid=norm_gene(r["gene_id"])
            if gid in out:raise ValueError("PRIMARY_DUPLICATE_ID "+gid)
            if not gid.startswith("ENSG"):raise ValueError("PRIMARY_NOT_ENSEMBL_ID")
            lfc=finite(r["logFC"]);fdr=finite(r["FDR"])
            if lfc is None or fdr is None:raise ValueError("PRIMARY_BAD_STAT")
            out[gid]={"gene_id":gid,"gene_symbol":r["gene_symbol"],
                      "log2fc":lfc,"fdr":fdr}
    if len(out)!=16050:raise ValueError("PRIMARY_EXPECTED_16050_GENES "+str(len(out)))
    return out

def load_candidates(path):
    with open(path,encoding="utf8",newline="") as h:
        r=csv.DictReader(h,delimiter="\t")
        if not {"gene_id","gene_symbol","literature_flag"}<=set(r.fieldnames or []):
            raise ValueError("P5G_FIELDS_MISSING")
        rows=list(r)
    if len(rows)!=86:raise ValueError("P5G_EXPECTED_86_CANDIDATES "+str(len(rows)))
    if len({norm_gene(r["gene_id"]) for r in rows})!=len(rows):
        raise ValueError("DUPLICATED_P5G_IDS")
    return rows

def correlation(xs,ys):
    if len(xs)!=len(ys) or len(xs)<20:raise ValueError("CORRELATION_NOT_DEFINED")
    xm=statistics.fmean(xs);ym=statistics.fmean(ys)
    cross=sum((a-xm)*(b-ym) for a,b in zip(xs,ys))
    xx=sum((a-xm)**2 for a in xs);yy=sum((b-ym)**2 for b in ys)
    return cross/math.sqrt(xx*yy)

def same_sign(a,b):return a*b>0

def summarize(primary,official,candidates):
    tm=official[TM_ALL]; high=official[TM_HIGH]
    shared=sorted(set(primary)&set(tm))
    paired=[(primary[g],tm[g]) for g in shared]
    comparable=[(a,b) for a,b in paired if b["fdr"] is not None]
    if len(shared)<9000:raise ValueError("LOW_SHARED_GENE_ID_OVERLAP")
    both_sig=[(a,b) for a,b in comparable if a["fdr"]<.05 and b["fdr"]<.05]
    original_de=[g for g in high if g in primary]
    high_replicated=[g for g in original_de
                     if primary[g]["fdr"]<.05 and same_sign(primary[g]["log2fc"],high[g]["log2fc"])]
    stratum=[(a,b) for a,b in comparable if a["fdr"]<.05 and abs(a["log2fc"])>=1]
    result={
      "status":"PASS_GENE_LEVEL_PUBLIC_TM_CROSSWALK",
      "study_primary":"GSE313736 high-confidence BeamA, 57 donor-paired edgeR",
      "study_external":"2026 IOVS independent cultured human TM, 48h DEX bulk",
      "source_official":SOURCE,"workbook":OFFICIAL_NAME,
      "source_primary_run":41085982185,
      "source_poolunit_run":41090827791,
      "primary_genes_tested":len(primary),
      "official_TM_all_rows":len(tm),
      "official_TM_high_confidence_rows":len(high),
      "shared_primary_official_gene_ids":len(shared),
      "external_missing_adjusted_p":sum(r["fdr"] is None for r in tm.values()),
      "primary_FDR_significant_shared":sum(primary[g]["fdr"]<.05 for g in shared),
      "external_FDR_significant_shared":sum(tm[g]["fdr"] is not None and tm[g]["fdr"]<.05 for g in shared),
      "FDR_significant_in_both":len(both_sig),
      "shared_log2fc_Pearson":round(correlation(
          [a["log2fc"] for a,b in paired],[b["log2fc"] for a,b in paired]),6),
      "all_shared_same_sign_n":sum(same_sign(a["log2fc"],b["log2fc"]) for a,b in paired),
      "all_shared_same_sign_fraction":round(sum(same_sign(a["log2fc"],b["log2fc"]) for a,b in paired)/len(paired),6),
      "both_FDR_significant_same_sign_n":sum(same_sign(a["log2fc"],b["log2fc"]) for a,b in both_sig),
      "both_FDR_significant_same_sign_fraction":round(
         sum(same_sign(a["log2fc"],b["log2fc"]) for a,b in both_sig)/max(1,len(both_sig)),6),
      "primary_strong_absLFC_ge1_n":len(stratum),
      "primary_strong_absLFC_ge1_same_sign_fraction":round(
         sum(same_sign(a["log2fc"],b["log2fc"]) for a,b in stratum)/max(1,len(stratum)),6),
      "external_high_conf_857_in_primary_table":len(original_de),
      "external_high_conf_reproduced_FDR_and_direction":len(high_replicated),
      "external_high_conf_reproduction_fraction":round(len(high_replicated)/max(1,len(original_de)),6),
      "comparability_caveats":[
          "Different donor sets, bulk TM versus BeamA cell-state pseudobulk; independent dose/protocol",
          "The 2026 IOVS experiment used 100 nM DEX; GSE313736 nominal 500 nM",
          "External FDR and primary FDR differ in gene universe and modeling method",
          "Direction agreement is not causal independence of technical treatment/library batches",
          "Not in supplementary table or not statistically significant there does NOT establish novelty",
          "Missing published adjusted p-values must not be interpreted as non-significance",
          "No direct steroid-induced clinical intraocular-pressure phenotype in primary cells"]
    }
    annotated=[]
    flags=Counter()
    for r in candidates:
        gid=norm_gene(r["gene_id"]); paper=tm.get(gid)
        paper_hi=gid in high
        a=primary[gid]
        if paper is None:
            status="NOT_IN_PUBLISHED_TM_ALL_TESTS"
        elif paper_hi:
            status="ALREADY_IN_PUBLISHED_857_HIGH_CONF_TM"
        elif paper["fdr"] is None:
            status="EXTERNAL_TM_ADJUSTED_P_UNAVAILABLE"
        elif paper["fdr"]<.05:
            status="ALREADY_PUBLISHED_TM_FDR_SIGNIFICANT"
        else:
            status="EXTERNAL_TM_TESTED_NOT_SIGNIFICANT"
        flags[status]+=1
        annotated.append({
           "gene_id":gid,"gene_symbol":r["gene_symbol"],
           "primary_log2fc":float(r["primary_log2fc"]),
           "primary_fdr":float(r["primary_fdr"]),
           "poolunit_log2fc":float(r["poolunit_log2fc"]),
           "manual_prior_literature_flag":r["literature_flag"],
           "official_TM_status":status,
           "official_TM_log2fc":paper["log2fc"] if paper else "",
           "official_TM_FDR":paper["fdr"] if paper and paper["fdr"] is not None else "",
           "official_TM_same_sign":same_sign(a["log2fc"],paper["log2fc"]) if paper else "",
           "novelty_not_verified":True})
    result["p5g_candidates_checked"]=len(annotated)
    result["p5g_official_TM_status_counts"]=dict(flags)
    result["p5g_in_official_high_conf"]=flags["ALREADY_IN_PUBLISHED_857_HIGH_CONF_TM"]
    result["p5g_not_in_official_all_TM_tested"]=flags["NOT_IN_PUBLISHED_TM_ALL_TESTS"]
    # Full lookup output is only derived comparisons on our own primary gene set,
    # not verbatim redistribution of the independent supplementary workbook.
    return result,annotated

def main():
    arg=argparse.ArgumentParser()
    arg.add_argument("--primary",required=True)
    arg.add_argument("--candidates",required=True)
    arg.add_argument("--output-dir",required=True)
    args=arg.parse_args()
    dest=Path(args.output_dir);dest.mkdir(parents=True,exist_ok=True)
    report={"status":"HOLD","created_utc":datetime.now(timezone.utc).isoformat()}
    try:
        zip_path=dest/"_temporary_original_public_supplement.zip"
        transferred=download_zip(zip_path)
        official=parse_official(zip_path)
        zip_path.unlink(missing_ok=True)
        base=load_primary(args.primary)
        candidates=load_candidates(args.candidates)
        report,selected=summarize(base,official,candidates)
        report["download_bytes"]=transferred
        with (dest/"P5j_86_candidates_external_TM_crosswalk.tsv").open("w",newline="",encoding="utf8") as h:
            writer=csv.DictWriter(h,fieldnames=list(selected[0]),delimiter="\t")
            writer.writeheader();writer.writerows(selected)
        print("P5J_OFFICIAL_TM_OVERLAP",report["shared_primary_official_gene_ids"],
              "both_FDR",report["FDR_significant_in_both"],
              "high_conf_replicated",report["external_high_conf_reproduced_FDR_and_direction"],
              "of",report["external_high_conf_857_in_primary_table"],
              "p5g_status",report["p5g_official_TM_status_counts"],flush=True)
    except Exception as e:
        report={"status":"HOLD_P5J_CROSSWALK_ERROR","error":type(e).__name__+": "+str(e)[:300]}
    finally:
        (dest/"_temporary_original_public_supplement.zip").unlink(missing_ok=True)
    (dest/"P5j_official_TM_crosswalk_summary.json").write_text(
        json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf8")
    if report["status"]!="PASS_GENE_LEVEL_PUBLIC_TM_CROSSWALK":
        raise SystemExit(2)

if __name__=="__main__":main()
