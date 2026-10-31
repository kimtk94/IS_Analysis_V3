#!/usr/bin/env python3
"""P7: independent TM EMT gene audit + ENA low-biomass/sample/control + shotgun manifest gates.
All analyses are public metadata or aggregate gene-level statistics; never raw patient reads.
"""
import argparse,csv,gzip,io,json,math,re,hashlib,urllib.parse,urllib.request
from collections import Counter
from pathlib import Path

TARGETS={"ERFE","CHAC1","PARM1","ADAMTS15","EGR3","LPAR1","ABHD3","DDIT4L","TFPI","STK17B","RAB32","SDC1"}
MARKERS={"MYOC","TGFBI","FN1","COL1A1","COL4A1","COL4A2","MMP2","TIMP1","THBS1","CCN2","CTGF","ACTA2"}
STUDIES={"PRJNA1145861":"conjunctival_16S","PRJEB74248":"conjunctival_shotgun","PRJNA1416399":"oral_16S"}
URL="https://www.ebi.ac.uk/ena/portal/api/filereport"
def get(url,limit=8_000_000):
    req=urllib.request.Request(url,headers={"User-Agent":"ocular-p7-research-provenance/1.0"})
    with urllib.request.urlopen(req,timeout=110) as h: data=h.read(limit+1)
    if len(data)>limit: raise ValueError("REMOTE_DATA_EXCEEDS_METADATA_LIMIT")
    return data
def enaq(accession,fields):
    q=urllib.parse.urlencode({"accession":accession,"result":"read_run","fields":",".join(fields),"format":"tsv"})
    raw=get(URL+"?"+q)
    rd=csv.DictReader(io.StringIO(raw.decode("utf8","replace")),delimiter="\t")
    if not set(fields).issubset(set(rd.fieldnames or [])):
        raise ValueError("ENA_MISSING_FIELDS "+accession+" "+str(rd.fieldnames))
    rows=list(rd)
    if not rows:raise ValueError("ENA_EMPTY_STUDY "+accession)
    if len({x["run_accession"] for x in rows})!=len(rows):
        raise ValueError("DUPLICATE_ENA_RUNS "+accession)
    return rows,hashlib.sha256(raw).hexdigest()
def num(s):
    try:
        x=float(s)
        return x if math.isfinite(x) else None
    except (TypeError,ValueError):
        return None
def writejson(out,name,v):
    (out/name).write_text(json.dumps(v,indent=2,ensure_ascii=False,allow_nan=False)+"\n",encoding="utf8")
def writetsv(out,name,records):
    if not records:raise ValueError("NO_TSV_ROWS")
    with (out/name).open("w",newline="",encoding="utf8") as f:
        wr=csv.DictWriter(f,fieldnames=list(records[0]),delimiter="\t")
        wr.writeheader();wr.writerows(records)
def load_primary(path):
    with gzip.open(path,"rt",encoding="utf8",newline="") as f: data=list(csv.DictReader(f,delimiter="\t"))
    if len(data)!=16050:raise ValueError("P5B_EXPECTED_16050")
    universe={}
    for r in data:
        gene=r["gene_id"].split(".")[0]
        v=num(r["logFC"]);fdr=num(r["FDR"])
        if not gene.startswith("ENSG") or v is None or fdr is None:
            raise ValueError("MISSING_PRIMARY_ENSG_OR_LOGFC_FDR")
        if gene in universe:raise ValueError("DUP_ENSEMBL_PRIMARY")
        universe[gene]={"gene_id":gene,"gene_symbol":r["gene_symbol"],
                        "primary_fc":v,"primary_fdr":fdr}
    return universe
def emt(primary,out):
    from openpyxl import load_workbook
    from ocular_p5m_timecourse_and_tiers import fetch_original_xlsx
    from ocular_p6d_crossstudy_pathway_concordance import GMT_SHA256,GMT_URL,pathway_es
    from urllib.request import urlopen
    src=load_primary(primary)
    official,nbytes=fetch_original_xlsx()
    wb=load_workbook(io.BytesIO(official),read_only=True,data_only=True)
    it=wb["SuppTable4_TM2Days_All_DEGs"].iter_rows(values_only=True)
    head=[str(x).strip() if x is not None else "" for x in next(it)]
    if not {"ID","log2FoldChange","padj"}.issubset(head):
        raise ValueError("EXTERNAL_OFFICIAL_SCHEMA_DRIFT")
    ix={k:head.index(k) for k in ("ID","log2FoldChange","padj")}
    matched={}
    n_external=0
    for row in it:
        if row[ix["ID"]] is None:continue
        n_external+=1
        gid=str(row[ix["ID"]]).strip().split(".")[0]
        if gid not in src:continue
        fc=num(row[ix["log2FoldChange"]])
        if fc is None:continue
        if gid in matched:raise ValueError("DUPLICATE_ENSG_EXTERNAL")
        matched[gid]={**src[gid],"external_fc":fc,"external_fdr":num(row[ix["padj"]])}
    wb.close()
    if n_external!=13313 or len(matched)!=12251:raise ValueError("P6D_48H_CROSS_STUDY_MATCH_DRIFT")
    ranked={}
    for x in matched.values():
        s=x["gene_symbol"]
        if s not in ranked or (abs(x["primary_fc"]),x["gene_id"])>(abs(ranked[s]["primary_fc"]),ranked[s]["gene_id"]):
            ranked[s]=x
    if len(ranked)!=12250:raise ValueError("P6D_SHARED_SYMBOL_UNIVERSE_DRIFT")
    dat=get(GMT_URL,limit=150000)
    if hashlib.sha256(dat).hexdigest()!=GMT_SHA256:
        raise ValueError("MSIGDB_VERSION_OR_HASH_DRIFT")
    sets={}
    for l in dat.decode("utf8").splitlines():
        z=l.split("\t")
        if z[0] in ("HALLMARK_EPITHELIAL_MESENCHYMAL_TRANSITION","HALLMARK_UNFOLDED_PROTEIN_RESPONSE"):
            sets[z[0]]=set(z[2:])
    if len(sets)!=2:raise ValueError("HALLMARK_EMT_UPR_MISSING")
    emt=sets["HALLMARK_EPITHELIAL_MESENCHYMAL_TRANSITION"]
    upr=sets["HALLMARK_UNFOLDED_PROTEIN_RESPONSE"]
    shared=list(ranked.values())
    hit=[r for r in shared if r["gene_symbol"] in emt]
    non=[r for r in shared if r["gene_symbol"] not in emt]
    if len(hit)!=194:raise ValueError("EMT_194_EXPECTED_UNIVERSE_DRIFT")
    def agree(rows):return sum(z["primary_fc"]*z["external_fc"]>0 for z in rows)
    def both(rows):return [z for z in rows if z["primary_fdr"]<.05 and z["external_fdr"] is not None and z["external_fdr"]<.05]
    if agree(hit)!=141 or len(both(hit))!=49:
        raise ValueError("P6D_EXPECTED_EMT_REPRODUCIBILITY_DRIFT")
    omits=[0,5,10,20]
    leave=[]
    for pathway,members in [("EMT",emt),("UPR",upr)]:
        candidate=sorted([r for r in shared if r["gene_symbol"] in members],
                         key=lambda r:(-max(abs(r["primary_fc"]),abs(r["external_fc"])),r["gene_id"]))
        for drop in omits:
            remove={r["gene_id"] for r in candidate[:drop]}
            common=[r for r in shared if r["gene_id"] not in remove]
            eps=[]
            for name,col in (("primary","primary_fc"),("external","external_fc")):
                inp=[{"gene_id":r["gene_id"],"gene_symbol":r["gene_symbol"],"lfc":r[col]} for r in common]
                eps.append(pathway_es(inp,members))
            leave.append({"pathway":pathway,"exclude_largest_max_abs_lfc_n":drop,
              "source_background_n":len(common),
              "primary_ES":round(eps[0],7),"external_ES":round(eps[1],7),
              "ES_signed_concordance":eps[0]*eps[1]>0,
              "filter_is_posthoc_robustness_not_preregistered":True})
    writetsv(out,"P7a_emt_upr_leave_top_effect_sensitivity.tsv",leave)
    hit.sort(key=lambda z:(z["primary_fc"]*z["external_fc"]<=0,
                           -min(abs(z["primary_fc"]),abs(z["external_fc"])),z["gene_id"]))
    genes=[]
    for x in hit:
        genes.append({"gene_id":x["gene_id"],"gene_symbol":x["gene_symbol"],
          "primary_500nM_log2fc":round(x["primary_fc"],6),
          "primary_FDR":x["primary_fdr"],"external_100nM_log2fc":round(x["external_fc"],6),
          "external_FDR":x["external_fdr"] if x["external_fdr"] is not None else "",
          "sign_agrees":x["primary_fc"]*x["external_fc"]>0,
          "both_study_FDR_lt_0_05":x["primary_fdr"]<.05 and x["external_fdr"] is not None and x["external_fdr"]<.05,
          "in_original_12_targets":x["gene_symbol"] in TARGETS})
    writetsv(out,"P7a_194_emt_gene_concordance.tsv",genes)
    marker_rows=[]
    for symbol in sorted(MARKERS|TARGETS):
        a=ranked.get(symbol)
        marker_rows.append({"symbol":symbol,"present_both":a is not None,
          "primary_lfc":round(a["primary_fc"],6) if a else "",
          "external_lfc":round(a["external_fc"],6) if a else "",
          "primary_FDR":a["primary_fdr"] if a else "",
          "external_FDR":a["external_fdr"] if a and a["external_fdr"] is not None else "",
          "same_signed_effect":a["primary_fc"]*a["external_fc"]>0 if a else "",
          "EMT_member":symbol in emt,"UPR_member":symbol in upr,
          "steroid_GC_TM_function_not_verified_by_this_test":True})
    writetsv(out,"P7a_12_targets_and_TM_ECM_marker_panel.tsv",marker_rows)
    source={
      "status":"PASS_CROSS_STUDY_EMT_UPr_GENE_194_EXPLORATORY",
      "official_100nM_bulk_48h_genes":n_external,
      "matched_ensembl_genes":len(matched),
      "matched_unique_symbols":len(shared),"EMT_genes":len(hit),
      "EMT_same_sign_count":agree(hit),"EMT_both_study_FDR_lt_0_05_count":len(both(hit)),
      "EMT_both_significant_opposite_count":sum(z["primary_fc"]*z["external_fc"]<0 for z in both(hit)),
      "background_nonEMT_same_sign_count":agree(non),
      "background_nonEMT_n":len(non),
      "EMT_and_background_depend_on_genecorrelation":True,
      "EMT_full_and_leave_top20_same_ES_direction":all(v["ES_signed_concordance"] for v in leave if v["pathway"]=="EMT"),
      "UPR_full_and_leave_top20_same_ES_direction":all(v["ES_signed_concordance"] for v in leave if v["pathway"]=="UPR"),
      "original_public_excel_download_bytes":nbytes,
      "no_pathway_p_values_or_cross_study_meta_analysis":True,
      "source_libraries_confound_treatment":"HOLD",
      "dose_100_vs_500_causal_comparison":"HOLD",
      "patient_steroid_delta_IOP":"STOP",
      "note":"Matched-gene signed agreement in two distinct TM studies is hypothesis evidence only. Both populations/platforms/culture doses differ. Leave-strong-gene-out is post-hoc sensitivity and not a new replication."}
    writejson(out,"P7a_emt_gene_sign_and_robustness.json",source)
    print("P7A",json.dumps(source,ensure_ascii=False),flush=True)

def samplemeta(out):
    fields=["run_accession","sample_accession","sample_title","sample_description",
            "library_strategy","library_layout"]
    fallback=["run_accession","sample_accession","sample_title","library_strategy","library_layout"]
    results=[]
    for study,label in STUDIES.items():
        error=None
        try:
            try:rows,sha=enaq(study,fields);used=fields
            except Exception as e:
                error=type(e).__name__+": "+str(e)[:130]
                rows,sha=enaq(study,fallback);used=fallback
            n=len(rows)
            case,healthy,unknown,ambiguous=0,0,0,0
            explicit_blanks=0;distinct_sample=set()
            for r in rows:
                distinct_sample.add(r["sample_accession"])
                t=((r.get("sample_title") or "")+" "+(r.get("sample_description") or "")).lower()
                # Keep controls separate from HEALTHY control; healthy-control ≠ extraction blank.
                isblank=bool(re.search(r"\b(?:extraction\s*blank|negative\s*control|reagent\s*blank|pcr\s*blank|no[\s-]*template\s*control|sampling\s*blank)\b",t))
                explicit_blanks+=int(isblank)
                diseased=bool(re.search(r"\b(?:glaucoma|poag|ocular[\s-]*hypertension|dry[\s-]*eye)\b",t))
                control=bool(re.search(r"\b(?:healthy|non[\s-]*glaucoma|normal[\s-]*control)\b",t))
                if diseased and control:ambiguous+=1
                elif diseased:case+=1
                elif control:healthy+=1
                else:unknown+=1
            result={"study":study,"scope":label,
               "status":"PASS_ENA_SAMPLE_METADATA_FIELD_AUDIT",
               "run_count":n,"unique_biosamples":len(distinct_sample),
               "fields_returned":used,"full_field_request_exception":error,
               "explicit_blank_control_run_mentions":explicit_blanks,
               "provisional_disease_keyword_runs":case,
               "provisional_healthy_keyword_runs":healthy,
               "ambiguous_control_and_disease_mentions":ambiguous,
               "unknown_or_nonclassifiable_runs":unknown,
               "valid_patient_group_assignment_from_title_alone":False,
               "actual_negative_blank_reads_verified":False,
               "glaucoma_eye_drops_or_steroid_drug_covariates_verified":False,
               "paired_patients_or_bilateral_eye_structure_verified":False,
               "clinical_steroid_dose_IOP_response_verified":False,
               "source_sha256":sha,
               "decision":"HOLD_LOW_BIOMASS_CONTROL_AND_CASE_CONTROL_COVARIATES"}
            results.append(result)
        except Exception as e:
            results.append({"study":study,"scope":label,
              "status":"HOLD_ENA_SAMPLE_METADATA_FETCH","error":type(e).__name__+": "+str(e)[:200],
              "valid_patient_group_assignment_from_title_alone":False})
    writejson(out,"P7b_three_microbiome_study_sample_design_gate.json",{
        "status":"PASS_METADATA_AUDIT_COMPLETED_MIXED_STUDY_HOLDS",
        "studies":results,
        "sample_identity_data_not_uploaded":True,
        "sample_title_keyword_classification_only_provisional":True,
        "published_methods_report_blanks_does_not_imply_raw_blank_fastqs_public":True,
        "causal_microbiome_role_in_steroid_glaucoma":"STOP"})
    print("P7B",[(x["study"],x["status"],x.get("run_count")) for x in results],flush=True)

def formats(out):
    accession="PRJEB74248"
    fields=["run_accession","library_layout","library_strategy","fastq_ftp",
            "submitted_ftp","submitted_md5","submitted_bytes"]
    rows,sha=enaq(accession,fields)
    if len(rows)!=62:raise ValueError("P6E_EXPECTED_62_SHOTGUN_RUNS")
    type_counts=Counter();link_counts=Counter();unknown_ext=Counter()
    nbytes=0;md5_present=0;submitted_file_count=0;unrecognized=0
    for r in rows:
        raw=(r["submitted_ftp"] or "").split(";")
        hashes=(r.get("submitted_md5") or "").split(";")
        sizes=(r.get("submitted_bytes") or "").split(";")
        if not raw or not raw[0]:raise ValueError("P6E_62_SUBMITTED_LINKS_DRIFT")
        hasfastq=bool(r["fastq_ftp"].strip())
        link_counts["fastq_run"]+=int(hasfastq)
        link_counts["submitted_only_run"]+=int(not hasfastq)
        for i,f in enumerate(raw):
            submitted_file_count+=1
            fname=f.rsplit("/",1)[-1].lower().split("?",1)[0]
            kind=next((suffix for suffix in (".fastq.gz",".fq.gz",".bam",".cram",".fastq",".fq",".sra") if fname.endswith(suffix)),"OTHER")
            type_counts[kind]+=1
            if kind=="OTHER":
                unknown_ext[fname.split(".")[-1]]+=1
                unrecognized+=1
            if len(hashes)>i and len(hashes[i])==32:md5_present+=1
            if len(sizes)>i:
                v=num(sizes[i])
                if v is not None and v>=0:nbytes+=int(v)
    result={"status":"PASS_ENA_SUBMITTED_FILE_FORMAT_MANIFEST_AUDIT",
        "accession":accession,"run_count":len(rows),
        "fastq_run_count":link_counts["fastq_run"],
        "submitted_only_run_count":link_counts["submitted_only_run"],
        "submitted_file_count":submitted_file_count,
        "file_extension_counts":dict(type_counts),
        "unrecognized_file_extension_counts":dict(unknown_ext),
        "md5_checksum_entry_count":md5_present,
        "sum_of_reported_submitted_bytes":nbytes,
        "ena_raw_metadata_sha256":sha,
        "http_head_or_actual_file_download_performed":False,
        "host_depletion_completed":False,
        "input_file_format_verified_by_binary_read":False,
        "clinical_case_control_and_negative_blanks":"HOLD",
        "interpretation":"Filename extension, submitted byte and MD5 metadata are not proof of readable FASTQ or verified raw download; further format and checksum preflight is required."}
    writejson(out,"P7c_shotgun_submitted_format_gate.json",result)
    print("P7C",json.dumps(result,ensure_ascii=False),flush=True)
def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--mode",choices=["emt","samplemeta","formats"],required=True)
    ap.add_argument("--primary",type=Path)
    ap.add_argument("--out",type=Path,required=True)
    a=ap.parse_args();a.out.mkdir(parents=True,exist_ok=True)
    if a.mode=="emt":
        if not a.primary:ap.error("emt requires --primary")
        emt(a.primary,a.out)
    elif a.mode=="samplemeta":samplemeta(a.out)
    else:formats(a.out)
if __name__=="__main__":
    main()
