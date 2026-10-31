#!/usr/bin/env python3
"""P7d official publication supplement / ENA group crosswalk discovery.

Reads source-supplied metadata, publishes ONLY counts and non-identifying supplement
filenames; no patient table, raw run accessions, personal IDs, or sample labels.
"""
import argparse,io,csv,json,re,urllib.request,urllib.parse,zipfile,hashlib
from pathlib import Path

URL="https://www.ebi.ac.uk/europepmc/webservices/rest/PMC11172891/supplementaryFiles"
PROJECT="PRJEB74248"
def fetch(url,limit=20_000_000):
    req=urllib.request.Request(url,headers={"User-Agent":"ocular-p7d-published-supplement-crosswalk-check/1.0"})
    with urllib.request.urlopen(req,timeout=120) as resp:
        data=resp.read(limit+1)
    if len(data)>limit:raise ValueError("SUPPLEMENT_EXCEEDS_SIZE_CAP")
    return data
def main():
    ap=argparse.ArgumentParser();ap.add_argument("--out",type=Path,required=True)
    args=ap.parse_args();args.out.mkdir(parents=True,exist_ok=True)
    url="https://www.ebi.ac.uk/ena/portal/api/filereport?"+urllib.parse.urlencode({
        "accession":PROJECT,"result":"read_run","fields":"run_accession,sample_accession,sample_alias,sample_title,experiment_alias,experiment_title","format":"tsv"})
    raw=fetch(url,limit=1_000_000)
    reader=csv.DictReader(io.StringIO(raw.decode("utf8","replace")),delimiter="\t")
    fields=set(reader.fieldnames or [])
    if not {"run_accession","sample_accession","sample_alias"}.issubset(fields):
        raise ValueError("ENA_ENRICHED_ALIAS_FIELDS_UNAVAILABLE")
    data=list(reader)
    if len(data)!=62 or len({r["run_accession"] for r in data})!=62:raise ValueError("EXPECTED_62_ENA_RUNS")
    known_runs={r["run_accession"] for r in data}
    known_samples={r["sample_accession"] for r in data}
    alias_case=0;alias_healthy=0;alias_ambig=0;alias_unknown=0
    for r in data:
        # Avoid experiment title, which often shares study-wide "glaucoma vs controls".
        s=((r.get("sample_alias") or "")+" "+(r.get("sample_title") or "")).lower()
        case=bool(re.search(r"\b(?:glaucoma|poag|pexg|patient)\b",s))
        healthy=bool(re.search(r"\b(?:healthy|normal|control)\b",s))
        if case and healthy:alias_ambig+=1
        elif case:alias_case+=1
        elif healthy:alias_healthy+=1
        else:alias_unknown+=1
    archive=fetch(URL)
    if not zipfile.is_zipfile(io.BytesIO(archive)):
        raise ValueError("EUROPE_PMC_SUPPLEMENT_NOT_ZIP")
    public_files=[];matches=set();content_scan_count=0
    def examine(z,prefix="",depth=0):
        nonlocal content_scan_count
        for name in z.namelist():
            if name.endswith("/"):continue
            entry=z.getinfo(name)
            if entry.file_size>10_000_000:continue
            data=z.read(name)
            leaf=name.rsplit("/",1)[-1]
            if leaf.lower().endswith(".zip") and depth<1 and zipfile.is_zipfile(io.BytesIO(data)):
                examine(zipfile.ZipFile(io.BytesIO(data)),prefix+leaf+"/",depth+1)
                continue
            kind=leaf.rsplit(".",1)[-1].lower() if "." in leaf else "unknown"
            rec={"name":(prefix+leaf)[:180],"extension":kind,"uncompressed_bytes":len(data)}
            texts=[]
            if kind in ("csv","tsv","txt"):
                texts.append(data.decode("utf8","replace"))
            elif kind=="xlsx":
                from openpyxl import load_workbook
                wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
                rec["worksheet_count"]=len(wb.sheetnames)
                for sheet in wb:
                    texts.append("\n".join("\t".join(str(c) for c in row if c is not None)
                           for row in sheet.iter_rows(values_only=True)))
                wb.close()
            for t in texts:
                content_scan_count+=1
                matches.update(known_runs.intersection(re.findall(r"\b(?:ERR|SRR|DRR)\d+\b",t)))
                matches.update(known_samples.intersection(re.findall(r"\b(?:ERS|SRS|DRS)\d+\b",t)))
            rec["parsed_tabular_content_for_accessions"]=bool(texts)
            public_files.append(rec)
    with zipfile.ZipFile(io.BytesIO(archive)) as z:
        examine(z)
    if not public_files:raise ValueError("SUPPLEMENT_ARCHIVE_NO_FILES")
    summary={"status":"PASS_PUBLIC_SUPPLEMENT_INVENTORY_HOLD_UNCONFIRMED_PATIENT_RUN_CROSSWALK",
      "bio_project":PROJECT,"source_publication_pmc":"PMC11172891",
      "europemc_supplement_size_bytes":len(archive),
      "supplement_archive_sha256":hashlib.sha256(archive).hexdigest(),
      "supplement_file_count":len(public_files),
      "supplement_files":public_files,
      "tabular_sections_scanned":content_scan_count,
      "ena_run_count":len(data),"ena_unique_samples":len(known_samples),
      "ena_sample_alias_case_keyword_count":alias_case,
      "ena_sample_alias_healthy_keyword_count":alias_healthy,
      "ena_sample_alias_ambiguous_keyword_count":alias_ambig,
      "ena_sample_alias_unclassified_count":alias_unknown,
      "exact_ena_sample_or_run_accessions_found_in_supplement_n":len(matches),
      "patient_to_run_mapping_verified":False,
      "negative_control_library_reads_validated":False,
      "steroid_exposure_and_iop_change_linked":False,
      "note":"Only exact run/sample accession string overlap was inspected; no published table has been used to impute individual clinical assignments. A clinically usable study design requires explicit verified crosswalk and medication/age confounding controls."}
    (args.out/"P7d_public_supplement_crosswalk_discovery.json").write_text(
      json.dumps(summary,indent=2,ensure_ascii=False)+"\n",encoding="utf8")
    print("P7D",json.dumps({k:summary[k] for k in ("status","supplement_file_count","tabular_sections_scanned","ena_sample_alias_case_keyword_count","ena_sample_alias_healthy_keyword_count","exact_ena_sample_or_run_accessions_found_in_supplement_n")},ensure_ascii=False),flush=True)
if __name__=="__main__":
    main()
