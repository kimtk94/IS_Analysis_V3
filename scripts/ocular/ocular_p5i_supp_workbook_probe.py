#!/usr/bin/env python3
"""P5i: inspect actual supplementary workbook from Europe PMC, metadata-only.

Uses the official PMC12882097 xlink file reference obtained from its XML.
No guessed full DEG sheet or automatic treatment assignment.
"""
import argparse,io,json,time,urllib.request,zipfile
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import load_workbook
SOURCE="https://www.ebi.ac.uk/europepmc/webservices/rest/PMC12882097/supplementaryFiles"
EXPECTED="iovs-67-2-13_s003.xlsx"
def get():
    last=None
    for i in range(3):
        try:
            req=urllib.request.Request(SOURCE,headers={"User-Agent":"ocular-public-study-supplement-audit/1.0"})
            with urllib.request.urlopen(req,timeout=60) as resp:
                b=resp.read(60*1024*1024+1)
            if len(b)>60*1024*1024:raise ValueError("ARCHIVE_GT_60MB")
            if not b.startswith(b"PK"):raise ValueError("NOT_ZIP")
            return b
        except Exception as e:
            last=e
            if i<2:time.sleep(i+1)
    raise last
def main():
    ap=argparse.ArgumentParser();ap.add_argument("--out",required=True);args=ap.parse_args()
    result={"source":SOURCE,"expected_original_supplement":EXPECTED,
            "study":"PMCID PMC12882097 / PMID 41636427", 
            "utc":datetime.now(timezone.utc).isoformat(),
            "status":"HOLD","preview_policy":"No entire gene tables in CI artifacts yet"}
    try:
        raw=get()
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            paths=[n for n in z.namelist() if not n.endswith("/")]
            result["archive_files"]=[x.split("/")[-1] for x in paths]
            xls=[n for n in paths if n.split("/")[-1].lower()==EXPECTED.lower()]
            if len(xls)!=1:raise ValueError("EXPECTED_OFFICIAL_WORKBOOK_NOT_FOUND")
            workbook_bytes=z.read(xls[0])
        result["original_xlsx_bytes"]=len(workbook_bytes)
        x=load_workbook(io.BytesIO(workbook_bytes),read_only=True,data_only=True)
        result["sheets"]=[]
        for sh in x:
            example=[]
            for i,row in enumerate(sh.iter_rows(min_row=1,max_row=min(6,sh.max_row or 6),values_only=True)):
                example.append([str(v)[:95] if v is not None else "" for v in row[:12]])
            result["sheets"].append({"sheet_name":sh.title,"rows":sh.max_row,
               "columns":sh.max_column,"header_preview":example})
        x.close()
        if not result["sheets"]:raise ValueError("XLSX_NO_SHEETS")
        result["status"]="PASS_OFFICIAL_SUPPLEMENTARY_XLSX_INSPECTED"
    except Exception as e:
        result["error"]=type(e).__name__+": "+str(e)[:280]
    p=Path(args.out);p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n")
    print("P5I",result["status"],"sheets",[s["sheet_name"] for s in result.get("sheets",[])],
          "error",result.get("error"))
    if result["status"]!="PASS_OFFICIAL_SUPPLEMENTARY_XLSX_INSPECTED":raise SystemExit(2)
if __name__=="__main__":main()
