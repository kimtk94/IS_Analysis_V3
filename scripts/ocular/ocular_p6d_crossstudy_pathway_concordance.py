#!/usr/bin/env python3
"""P6d cross-study pathway direction check: matched Ensembl genes, 500nM BeamA vs official 100nM 48h bulk TM.

Compare pathways and source FDR with shared mapped gene background. Do not
interpret these cohort/cell/dose/protocol-confounded effects as a dose-response.
"""
import argparse,csv,gzip,io,json,math,hashlib
from pathlib import Path
from collections import Counter
import numpy as np
from openpyxl import load_workbook
from ocular_p5m_timecourse_and_tiers import fetch_original_xlsx, XLSX

GMT_URL="https://data.broadinstitute.org/gsea-msigdb/msigdb/release/2025.1.Hs/h.all.v2025.1.Hs.symbols.gmt"
GMT_SHA256="f22066af72e215ccb7b89d88e492c07e1eef17534c2ca7b0f9902cfecbbdd8e9"
SETS={
 "HALLMARK_TNFA_SIGNALING_VIA_NFKB","HALLMARK_INFLAMMATORY_RESPONSE",
 "HALLMARK_TGF_BETA_SIGNALING","HALLMARK_REACTIVE_OXYGEN_SPECIES_PATHWAY",
 "HALLMARK_EPITHELIAL_MESENCHYMAL_TRANSITION","HALLMARK_IL6_JAK_STAT3_SIGNALING",
 "HALLMARK_UNFOLDED_PROTEIN_RESPONSE","HALLMARK_HYPOXIA","HALLMARK_APOPTOSIS"
}
TARGETS={"ERFE","CHAC1","PARM1","ADAMTS15","EGR3","LPAR1","ABHD3",
         "DDIT4L","TFPI","STK17B","RAB32","SDC1"}
def number(x):
    try:
        f=float(x)
        return f if math.isfinite(f) else None
    except (TypeError,ValueError):
        return None
def tsv(path,rows):
    if not rows:raise ValueError("EMPTY_REPORT")
    with path.open("w",encoding="utf8",newline="") as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]),delimiter="\t")
        writer.writeheader();writer.writerows(rows)
def pathway_es(data,member):
    # Data contain (gene_id,symbol,lfc) for the same matched background in each cohort.
    ordered=sorted(data,key=lambda r:(-r["lfc"],r["gene_id"]))
    values=np.asarray([abs(r["lfc"]) for r in ordered],dtype="float64")
    hits=np.asarray([r["gene_symbol"] in member for r in ordered],dtype=bool)
    nh=int(hits.sum());n=len(ordered)
    if nh<10 or nh>=n:raise ValueError("LOW_GENESET_OVERLAP")
    sw=values[hits].sum()
    if not sw>0:raise ValueError("ZERO_PATHWAY_WEIGHTS")
    c=np.cumsum(np.where(hits,values/sw,-1/(n-nh)))
    return float(c[int(np.argmax(np.abs(c)))])
def main():
    a=argparse.ArgumentParser()
    a.add_argument("--primary",required=True,type=Path)
    a.add_argument("--out",required=True,type=Path)
    args=a.parse_args()
    args.out.mkdir(parents=True,exist_ok=True)
    with gzip.open(args.primary,"rt",encoding="utf8",newline="") as f:
        source=list(csv.DictReader(f,delimiter="\t"))
    if len(source)!=16050:raise ValueError("PRIMARY_SOURCE_DRIFT")
    byid={}
    for r in source:
        gid=r["gene_id"].split(".")[0]
        fc=number(r["logFC"]);fdr=number(r["FDR"])
        if not gid.startswith("ENSG") or fc is None or fdr is None:
            raise ValueError("PRIMARY_GENE_ID_OR_NUMERIC_DRIFT")
        if gid in byid:raise ValueError("DUP_PRIMARY_ENSEMBL_ID")
        byid[gid]={"gene_id":gid,"gene_symbol":r["gene_symbol"],"lfc":fc,"fdr":fdr}
    original_zip,archive_len=fetch_original_xlsx()
    wb=load_workbook(io.BytesIO(original_zip),read_only=True,data_only=True)
    table=wb["SuppTable4_TM2Days_All_DEGs"]
    itr=table.iter_rows(values_only=True)
    header=[str(x).strip() if x is not None else "" for x in next(itr)]
    needed={"ID","log2FoldChange","padj"}
    if not needed.issubset(header):raise ValueError("EXTERNAL_OFFICIAL_SHEET_FIELDS_CHANGED")
    inds={x:header.index(x) for x in needed}
    matched=[]
    external_rows=0
    duplicated=0
    seen=set()
    for cells in itr:
        if cells[inds["ID"]] is None:continue
        external_rows+=1
        gid=str(cells[inds["ID"]]).split(".")[0].strip()
        if gid not in byid:continue
        if gid in seen:
            duplicated+=1
            continue
        seen.add(gid)
        fc=number(cells[inds["log2FoldChange"]])
        efdr=number(cells[inds["padj"]])
        if fc is None:continue
        x=byid[gid]
        matched.append({"gene_id":gid,"gene_symbol":x["gene_symbol"],
          "primary_lfc":x["lfc"],"primary_fdr":x["fdr"],
          "external_lfc":fc,"external_fdr":efdr})
    wb.close()
    if external_rows!=13313 or duplicated:
        raise ValueError("OFFICIAL_48H_GENE_ROWS_OR_DUPLICATES_CHANGED")
    if len(matched)<10000:raise ValueError("INSUFFICIENT_MATCHED_ENSEMBL_UNIVERSE "+str(len(matched)))
    with __import__("urllib.request",fromlist=["urlopen"]).urlopen(GMT_URL,timeout=70) as resp:
        data=resp.read(150001)
    if len(data)>150000 or hashlib.sha256(data).hexdigest()!=GMT_SHA256:
        raise ValueError("HALLMARK_GMT_VERSION_OR_SHA256_CHANGED")
    sets={}
    for line in data.decode().splitlines():
        cols=line.split("\t")
        if cols[0] in SETS:sets[cols[0]]=set(cols[2:])
    if set(sets)!=SETS:raise ValueError("PATHWAY_SET_MISSING")
    # Duplicate symbols in same Ensembl reference background: deterministically choose
    # primary most abs(FC) row, keeping that same Ensembl ID in BOTH studies.
    dedup={}
    for x in matched:
        symbol=x["gene_symbol"]
        if symbol not in dedup or (abs(x["primary_lfc"]),x["gene_id"])>(abs(dedup[symbol]["primary_lfc"]),dedup[symbol]["gene_id"]):
            dedup[symbol]=x
    shared=list(dedup.values())
    rows=[]
    concordance={}
    for name in sorted(SETS):
        member=sets[name]
        hits=[r for r in shared if r["gene_symbol"] in member]
        primary=pathway_es([{"gene_id":r["gene_id"],"gene_symbol":r["gene_symbol"],"lfc":r["primary_lfc"]} for r in shared],member)
        external=pathway_es([{"gene_id":r["gene_id"],"gene_symbol":r["gene_symbol"],"lfc":r["external_lfc"]} for r in shared],member)
        agree=sum(r["primary_lfc"]*r["external_lfc"]>0 for r in hits)
        valid=[r for r in hits if r["external_fdr"] is not None]
        sig=[r for r in valid if r["primary_fdr"]<.05 and r["external_fdr"]<.05]
        rows.append({"hallmark":name,"matched_background_genes":len(shared),
          "hallmark_overlap_n":len(hits),
          "primary_500nM_BeamA_ES":round(primary,6),
          "published_100nM_bulk_48h_ES":round(external,6),
          "ES_direction_same":primary*external>0,
          "gene_level_logFC_sign_same_n":agree,
          "gene_level_logFC_sign_same_fraction":round(agree/len(hits),4),
          "both_cohorts_FDR_below_0_05_n":len(sig),
          "both_significant_opposite_n":sum(r["primary_lfc"]*r["external_lfc"]<0 for r in sig),
          "12_target_set_members":",".join(sorted(member&TARGETS)),
          "experiment_conditions_not_matched":True,
          "dose_reversal_established":False})
    tsv(args.out/"P6d_matched_background_9_hallmark_crossstudy.tsv",rows)
    candidates=[{"gene_symbol":r["gene_symbol"],"gene_id":r["gene_id"],
       "primary_500nM_log2fc":round(r["primary_lfc"],6),
       "external_100nM_log2fc":round(r["external_lfc"],6),
       "primary_FDR":r["primary_fdr"],
       "external_FDR":r["external_fdr"],
       "direction_matches":r["primary_lfc"]*r["external_lfc"]>0}
      for r in shared if r["gene_symbol"] in TARGETS]
    if len(candidates)!=12:raise ValueError("NOT_ALL_12_SHARED_TARGETS")
    tsv(args.out/"P6d_12_shared_target_48h_signs.tsv",sorted(candidates,key=lambda r:r["gene_symbol"]))
    if any(x["direction_matches"] for x in candidates):
        raise ValueError("ORIGINAL_12_OPPOSITE_SIGN_CONTRACT_DRIFT")
    result={
       "status":"PASS_48H_CROSS_STUDY_MATCHED_BACKGROUND_DESCRIPTIVE",
       "original_primary_filtered_genes":len(source),"external_100nM_48h_rows":external_rows,
       "matched_ensembl_ids":len(matched),"matched_unique_symbols":len(shared),
       "gmt_sha256":GMT_SHA256,"n_gene_sets":len(rows),
       "gene_sets_same_direction":sum(x["ES_direction_same"] for x in rows),
       "emt_sign_same":next(x["ES_direction_same"] for x in rows if x["hallmark"]=="HALLMARK_EPITHELIAL_MESENCHYMAL_TRANSITION"),
       "upr_sign_same":next(x["ES_direction_same"] for x in rows if x["hallmark"]=="HALLMARK_UNFOLDED_PROTEIN_RESPONSE"),
       "source_official":"Mehrotra et al. 2026 100nM bulk TM, 48h supplementary workbook",
       "source_archive_download_bytes":archive_len,
       "no_crosscohort_p_value_or_meta_analysis":True,
       "cross_cohort_same_dose_and_cell_type":False,
       "clinical_delta_IOP_reproducibility":"NOT_TESTED",
       "interpretation":"Pathway ES directional agreement is descriptive only. Bulk 100nM and BeamA 500nM differ by donor cohorts, cell definition, platform and protocols; no causal dose-specific interaction is testable."}
    (args.out/"P6d_crossstudy_gate_summary.json").write_text(
       json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf8")
    print("P6D",json.dumps(result,ensure_ascii=False),flush=True)
if __name__=="__main__":
    main()
