#!/usr/bin/env python3
"""Stage35 read-only original PDC Experimental Design XLSX bridge preflight.

Compare only in-memory exact source tokens and row co-occurrences against CDAP
.sample.txt and official PDC aliquot/run namespaces. No row order or fuzzy
identities. No XLSX cell contents, URLs, filenames, IDs or source rows exported.
"""
import argparse
import collections
import datetime as dt
import hashlib
import io
import json
import pathlib
import re

from openpyxl import load_workbook
from probe_stage30_gao_catalog import PDC, pdc_query
from probe_stage30_gao_quant_header import safe_public_https
from probe_stage30_gao_quant_api_header_to_biospecimen import official_biospecimen
from audit_stage33_original_cdap_sample_bridge import discover, read_bytes, parse_sample_table, RATIO_PATTERN
from audit_stage31_pdc_tmt_header_templates import run_design

MAX_SHEETS=20
MAX_ROWS=6000
MAX_COLS=250
MAX_CELLS=150000

def lookup_other_metadata(pid, provider=pdc_query):
    query=('{filesPerStudy(pdc_study_id:'+json.dumps(pid)+
           ',data_category:"Other Metadata",offset:0,limit:30)'
           '{file_name data_category file_size md5sum signedUrl{url}}}')
    obj, digest=provider(query)
    files=obj.get("filesPerStudy")
    if not isinstance(files,list) or len(files)>30 or any(
        not isinstance(x,dict) or x.get("data_category")!="Other Metadata" for x in files):
        raise ValueError("untrusted PDC other metadata listing")
    xlsx=[]
    for r in files:
        if not str(r.get("file_name") or "").lower().endswith(".xlsx"):
            continue
        links=r.get("signedUrl") or []
        if isinstance(links,dict):links=[links]
        url=next((u.get("url") for u in links if isinstance(u,dict) and safe_public_https(u.get("url"))),None)
        xlsx.append({"url":url,"size":r.get("file_size"),"md5":r.get("md5sum"),"gz":False})
    return xlsx,digest

def in_scope_tokens(rows,reporters):
    analytical=set()
    biological=set()
    ratio_count=0
    for row in rows:
        name=str(row.get("AnalyticalSample") or "").strip()
        if name: analytical.add(name)
        for spec in str(row.get("Ratios") or "").split(","):
            m=RATIO_PATTERN.fullmatch(spec.strip())
            if not m: continue
            reporter_code=m.group(1).upper()
            matched=[r for r in reporters if r.upper().replace("TMT_","").replace("TMT-","")==reporter_code]
            if len(matched)!=1: continue
            label=str(row.get(matched[0]) or "").strip()
            if label and label.upper()!="POOL":
                biological.add(label);ratio_count+=1
    return analytical,biological,ratio_count

def source_universe(aliquots,design):
    vals={"official_aliquot_uuid":set(aliquots["aliquot_id"]),
          "official_aliquot_submitter":set(aliquots["aliquot_submitter_id"]),
          "official_plex":{x["plex"] for x in design if x["plex"]},
          "official_run":{x["run"] for x in design if x["run"]}}
    return vals

def text_tokens(row):
    # Do not convert numeric XLSX cells into sample identifiers:
    # Excel may have already discarded leading zeros from apparent IDs.
    return {v.strip() for v in row if isinstance(v,str) and v.strip()}

def compare_worksheet(rows,source):
    """Return aggregated exact literal overlap and unique source row pairs."""
    sums=collections.Counter()
    role_hits=collections.Counter()
    link=collections.defaultdict(set)
    workbook_cols=0
    for row_idx, row in enumerate(rows):
        if row_idx>=MAX_ROWS:raise ValueError("sheet row bound exceeded")
        if len(row)>MAX_COLS:raise ValueError("sheet width bound exceeded")
        workbook_cols=max(workbook_cols,len(row))
        vals=text_tokens(row)
        # Header words indicate possible schema, not reliable field semantics.
        if row_idx < 10:
            for v in vals:
                for term in ("sample","aliquot","plex","run","channel","tmt"):
                    if term in v.casefold():role_hits[term]+=1
        hits={k: vals & allowed for k,allowed in source.items()}
        for k,matched in hits.items():
            sums[k+"_row_memberships"]+=len(matched)
        if len(hits["cdap_biological_label"])==1 and (
            len(hits["official_aliquot_uuid"]|hits["official_aliquot_submitter"])==1):
            a=next(iter(hits["cdap_biological_label"]))
            b=next(iter(hits["official_aliquot_uuid"]|hits["official_aliquot_submitter"]))
            link[a].add(b)
            sums["single_source_row_pair_candidates_n"]+=1
        if hits["cdap_biological_label"] and (
            hits["official_aliquot_uuid"] or hits["official_aliquot_submitter"]):
            sums["rows_with_both_label_and_official_ID_n"]+=1
            if len(hits["cdap_biological_label"])>1 or (
                len(hits["official_aliquot_uuid"]|hits["official_aliquot_submitter"])>1):
                sums["rows_with_ambiguous_multiple_link_tokens_n"]+=1
        sums["rows_scanned_n"]+=1
    singly={a:next(iter(v)) for a,v in link.items() if len(v)==1}
    # A co-occurrence is NOT a valid mapping until field semantics are verified.
    return {
        "rows_scanned_n":sums["rows_scanned_n"],"max_width_n":workbook_cols,
        "textual_schema_keyword_hits_top_rows":dict(role_hits),
        "source_identifier_row_memberships":{k:sums[k+"_row_memberships"] for k in source},
        "rows_with_both_label_and_official_ID_n":sums["rows_with_both_label_and_official_ID_n"],
        "rows_with_ambiguous_multiple_link_tokens_n":sums["rows_with_ambiguous_multiple_link_tokens_n"],
        "single_source_row_pair_candidates_n":sums["single_source_row_pair_candidates_n"],
        "distinct_source_labels_with_single_candidate_n":len(singly),
        "distinct_official_aliases_with_single_candidate_n":len(set(singly.values())),
        "row_cooccurrence_is_mapping_proof":False,
        "source_identifiers_and_cell_text_exported":False,
    }

def inspect_xlsx(blob,source):
    book=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    try:
        if len(book.worksheets)>MAX_SHEETS:
            raise ValueError("too many experimental design worksheets")
        all_sheets=[]
        scanned_cells=0
        for ws in book.worksheets:
            if ws.max_row and ws.max_row>MAX_ROWS:
                raise ValueError("experimental design sheet exceeds row budget")
            if ws.max_column and ws.max_column>MAX_COLS:
                raise ValueError("experimental design sheet exceeds column budget")
            stat=compare_worksheet(ws.iter_rows(values_only=True),source)
            scanned_cells+=stat["rows_scanned_n"]*stat["max_width_n"]
            if scanned_cells>MAX_CELLS:raise ValueError("too many experimental-design cells")
            all_sheets.append(stat)
        return {"worksheet_count_n":len(all_sheets),
                "total_scanned_cells_upper_bound_n":scanned_cells,
                "worksheet_anonymous_statistics":all_sheets,
                "all_sheets_fully_reviewed_within_bounds":True}
    finally:
        book.close()

def one(pid):
    out={"pdc_study_id":pid,"status":"ACCESS_HOLD",
         "file_names_urls_or_original_values_exported":False,
         "source_CDP_to_PDC_mapping_certified":False,
         "quant_header_to_aliquot_certified":False,
         "RNA_protein_patient_join_certified":False}
    try:
        files,catalog_digest=lookup_other_metadata(pid)
        out.update({"source_metadata_catalog_sha256":catalog_digest,
                    "official_xlsx_candidates_n":len(files)})
        if len(files)!=1:
            out["status"]="REVIEW_XLSX_CANDIDATE_COUNT_NOT_ONE";return out
        xlsx=read_bytes(files[0])
        out["original_xlsx_SHA256"]=hashlib.sha256(xlsx).hexdigest()
        companion,cdap_catalog_sha=discover(pid)
        out["CDAP_source_catalog_SHA256"]=cdap_catalog_sha
        if len(companion)!=1:
            out["status"]="REVIEW_CDAP_CANDIDATE_COUNT_NOT_ONE";return out
        cdap=read_bytes(companion[0])
        out["original_CDAP_sample_SHA256"]=hashlib.sha256(cdap).hexdigest()
        rows,reporters=parse_sample_table(cdap)
        analytical,biological,ratio_n=in_scope_tokens(rows,reporters)
        if len(rows)!=33 or len(biological)!=330 or ratio_n!=330:
            raise ValueError("original CDAP 330 label universe drift")
        design,ref,design_sha=run_design(pid)
        aliases=official_biospecimen(pid)
        source=source_universe(aliases,design)
        source.update({"cdap_analytical_sample":analytical,
                       "cdap_biological_label":biological})
        out.update({"source_official_experiment_design_SHA256":design_sha,
                    "source_analytical_sample_n":len(analytical),
                    "source_cdap_biological_labels_n":len(biological),
                    "source_nonreference_official_aliquots_n":len({x["aliquot"] for x in design if x["aliquot"] not in ref}),
                    "xlsx_source":inspect_xlsx(xlsx,source),
                    "status":"ORIGINAL_EXPERIMENTAL_XLSX_SCHEMA_AND_EXACT_TOKEN_OVERLAP_REVIEWED"})
    except Exception as e:
        out["error_class"]=type(e).__name__
    return out

def review(obj):
    if obj.get("schema")!="hcc-stage35-original-PDC-XLSX-bridge-v1":
        raise ValueError("bad schema")
    rows=obj.get("sources")
    if not isinstance(rows,list) or {x.get("pdc_study_id") for x in rows}!=set(PDC):
        raise ValueError("wrong PDC studies")
    for x in rows:
        for k in ("file_names_urls_or_original_values_exported","source_CDP_to_PDC_mapping_certified",
                  "quant_header_to_aliquot_certified","RNA_protein_patient_join_certified"):
            if x.get(k) is not False:raise ValueError("unverified biological identity or source disclosure")
        if x["status"]=="ORIGINAL_EXPERIMENTAL_XLSX_SCHEMA_AND_EXACT_TOKEN_OVERLAP_REVIEWED":
            s=x["xlsx_source"]
            if s["worksheet_count_n"]!=len(s["worksheet_anonymous_statistics"]):
                raise ValueError("worksheet count drift")
            for ws in s["worksheet_anonymous_statistics"]:
                if ws.get("row_cooccurrence_is_mapping_proof") is not False or ws.get("source_identifiers_and_cell_text_exported") is not False:
                    raise ValueError("unsupported identity claim")
        elif x["status"] not in ("ACCESS_HOLD","REVIEW_XLSX_CANDIDATE_COUNT_NOT_ONE",
                               "REVIEW_CDAP_CANDIDATE_COUNT_NOT_ONE"):
            raise ValueError("bad source status")
    return {"schema":"hcc-stage35-PDC-source-design-postflight-v1",
        "original_XLSX_source_schema_completed_n":sum(x["status"]=="ORIGINAL_EXPERIMENTAL_XLSX_SCHEMA_AND_EXACT_TOKEN_OVERLAP_REVIEWED" for x in rows),
        "G30_B":"HOLD_EXACT_330_QUANT_COLUMN_TO_BIOLOGICAL_ALIQUOT_UNPROVEN",
        "G30_C":"HOLD_NODE_RNA_ORIGINAL_SAMPLE_ID",
        "G30_D":"STOP_NO_MATCHED_EXTERNAL_PATIENT_MULTIMODAL_DATA",
        "biological_or_clinical_claims_allowed":False}

def main():
    p=argparse.ArgumentParser();p.add_argument("--out",required=True);p.add_argument("--review-out",required=True)
    a=p.parse_args()
    result={"schema":"hcc-stage35-original-PDC-XLSX-bridge-v1",
            "checked_utc":dt.datetime.now(dt.timezone.utc).isoformat(),
            "sources":[one(pid) for pid in PDC]}
    verdict=review(result)
    for filename,x in ((a.out,result),(a.review_out,verdict)):
        out=pathlib.Path(filename);out.parent.mkdir(parents=True,exist_ok=True)
        out.write_text(json.dumps(x,indent=2)+"\n")
    print("STAGE35_XLSX_SOURCE",json.dumps(result))
    print("STAGE35_REVIEW",json.dumps(verdict))
if __name__=="__main__":main()
