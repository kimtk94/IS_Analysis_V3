#!/usr/bin/env python3
"""
Lightweight Stage-0 input inventory for the KoGES Metabolic Resilience project.

Purpose
-------
Identify participant-level KoGES files and their schemas WITHOUT loading full
datasets into memory. This is deliberately conservative because large cohort
files can terminate low-memory sessions when read eagerly.

Usage
-----
python3 audit_metabolic_resilience_koges_inputs.py \
  --koges-root /srv/is-analysis/data/koges \
  --output-root /srv/is-analysis/results/metabolic_resilience/stage0_feasibility

Outputs
-------
KOGES_FILE_INVENTORY.tsv
KOGES_RAW_SCHEMA.tsv
STAGE0_INPUT_AUDIT.md
"""

from __future__ import annotations
import argparse
import csv
import os
import sys
from pathlib import Path

SPECIAL_MISSING = {"55555", "66666", "77777", "99999"}
ANCHORS = [
    "ID","DATA_CLASS","VISIT","EDATE","AGE","SEX",
    "HEIGHT","WEIGHT","WAIST",
    "SBP_L","DBP_L","SBP_R","DBP_R",
    "GLU0_ORI","GLU0_TR","HBA1C",
    "TRIGLY_ORI","TRIGLY_TR","HDL_ORI","HDL_TR",
    "HTN","DM","LIP",
    "TREATD1","TREATD2","TREATD9",
    "DRUGHT","DRUGDM","DRUGLP",
    "DRINK","TOTALC",
    "MET","EXERCUR","EXER","EXERFQ","EXERDU",
]

TEXT_EXT = {".csv", ".tsv", ".txt"}
PARQUET_EXT = {".parquet", ".pq"}
EXCEL_EXT = {".xlsx", ".xls"}
SAS_EXT = {".sas7bdat", ".xpt"}
SPSS_EXT = {".sav"}

def human_size(n: int) -> str:
    for u in ["B","KB","MB","GB","TB"]:
        if n < 1024 or u == "TB":
            return f"{n:.1f}{u}"
        n /= 1024

def sniff_delimiter(path: Path) -> str:
    if path.suffix.lower() == ".tsv":
        return "\t"
    try:
        with path.open("r", encoding="utf-8-sig", errors="replace") as f:
            sample = f.read(65536)
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except Exception:
        return ","

def schema_text(path: Path):
    delim = sniff_delimiter(path)
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.reader(f, delimiter=delim)
        header = next(reader, [])
    return header, f"delimiter={repr(delim)}"

def schema_parquet(path: Path):
    try:
        import pyarrow.parquet as pq
        pf = pq.ParquetFile(path)
        return pf.schema_arrow.names, f"row_groups={pf.num_row_groups}"
    except Exception as e:
        return [], f"ERROR:{type(e).__name__}:{e}"

def schema_excel(path: Path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        # Only inspect sheet names + first nonempty row from each sheet.
        best_cols = []
        notes = []
        for s in wb.sheetnames[:20]:
            ws = wb[s]
            cols = []
            for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
                vals = [str(x).strip() if x is not None else "" for x in row]
                nonblank = [x for x in vals if x]
                if len(nonblank) >= 3:
                    cols = vals
                    break
            notes.append(f"{s}:{len(cols)}")
            if len(cols) > len(best_cols):
                best_cols = cols
        return best_cols, "sheets=" + ",".join(notes)
    except Exception as e:
        return [], f"ERROR:{type(e).__name__}:{e}"

def schema_sas_or_spss(path: Path):
    # pyreadstat metadataonly=True avoids loading observations.
    try:
        import pyreadstat
        ext = path.suffix.lower()
        if ext == ".sas7bdat":
            _, meta = pyreadstat.read_sas7bdat(path, metadataonly=True)
        elif ext == ".xpt":
            _, meta = pyreadstat.read_xport(path, metadataonly=True)
        elif ext == ".sav":
            _, meta = pyreadstat.read_sav(path, metadataonly=True)
        else:
            return [], "unsupported"
        return list(meta.column_names), f"rows={getattr(meta,'number_rows',None)}"
    except Exception as e:
        return [], f"ERROR:{type(e).__name__}:{e}"

def get_schema(path: Path):
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        return schema_text(path)
    if ext in PARQUET_EXT:
        return schema_parquet(path)
    if ext in EXCEL_EXT:
        return schema_excel(path)
    if ext in SAS_EXT | SPSS_EXT:
        return schema_sas_or_spss(path)
    return [], "not_inspected"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--koges-root", required=True)
    ap.add_argument("--output-root", required=True)
    ap.add_argument("--max-schema-file-gb", type=float, default=20.0)
    args = ap.parse_args()

    root = Path(args.koges_root).expanduser().resolve()
    out = Path(args.output_root).expanduser().resolve()
    out.mkdir(parents=True, exist_ok=True)

    if not root.exists():
        raise SystemExit(f"[ERROR] KoGES root does not exist: {root}")

    inv_path = out / "KOGES_FILE_INVENTORY.tsv"
    sch_path = out / "KOGES_RAW_SCHEMA.tsv"
    md_path = out / "STAGE0_INPUT_AUDIT.md"

    files = sorted(p for p in root.rglob("*") if p.is_file())

    with inv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter="\t")
        w.writerow(["path","relative_path","suffix","size_bytes","size_human"])
        for p in files:
            st = p.stat()
            w.writerow([str(p), str(p.relative_to(root)), p.suffix.lower(), st.st_size, human_size(st.st_size)])

    schema_rows = []
    for i,p in enumerate(files,1):
        size_gb = p.stat().st_size / (1024**3)
        if size_gb > args.max_schema_file_gb:
            schema_rows.append([str(p.relative_to(root)), p.suffix.lower(), "", "", "SKIPPED_SIZE_GATE"])
            continue
        cols, note = get_schema(p)
        colset = {str(c).strip().upper() for c in cols if str(c).strip()}
        hits = [a for a in ANCHORS if a.upper() in colset]
        schema_rows.append([
            str(p.relative_to(root)),
            p.suffix.lower(),
            len(cols),
            ",".join(hits),
            note
        ])
        if i % 25 == 0:
            print(f"[audit] inspected {i}/{len(files)} files", file=sys.stderr)

    with sch_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter="\t")
        w.writerow(["relative_path","suffix","n_columns","anchor_hits","schema_note"])
        w.writerows(schema_rows)

    likely_raw = [r for r in schema_rows if r[3]]
    report = [
        "# KoGES Metabolic Resilience — Stage 0 Input Audit",
        "",
        f"- Root: `{root}`",
        f"- Files found: **{len(files)}**",
        f"- Files with at least one anchor-column hit: **{len(likely_raw)}**",
        "",
        "## Candidate participant-level files",
        "",
    ]
    if likely_raw:
        for r in likely_raw[:100]:
            report.append(f"- `{r[0]}` — anchors: `{r[3]}`")
    else:
        report += [
            "No file exposed exact integrated anchor column names during the lightweight schema scan.",
            "",
            "This does not prove that participant-level data are absent; wave-specific source columns may use different names.",
            "Next inspect the largest CSV/SAS/Parquet files and their codebooks without full-table loading.",
        ]
    report += [
        "",
        "## Safety note",
        "",
        "This audit intentionally avoids loading complete datasets into memory.",
        "Do not replace it with a recursive pandas `read_*` loop over every cohort file.",
    ]
    md_path.write_text("\n".join(report), encoding="utf-8")

    print(f"[OK] {inv_path}")
    print(f"[OK] {sch_path}")
    print(f"[OK] {md_path}")

if __name__ == "__main__":
    main()
